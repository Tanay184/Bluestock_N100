"""Peer-group percentile ranks, badges, radar charts and peer_comparison.xlsx."""

import re
import sqlite3

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

from src.analytics.screener import build_snapshot
from src.etl.loader import DB_PATH, OUT, ROOT

METRICS = [
    "roe",
    "roce",
    "npm",
    "opm",
    "de",
    "fcf_margin",
    "rev_cagr_5yr",
    "pat_cagr_5yr",
    "eps_cagr_5yr",
    "asset_turnover",
]
LOWER_BETTER = {"de"}
RADAR = [
    "roe",
    "roce",
    "npm",
    "de",
    "fcf_margin",
    "pat_cagr_5yr",
    "rev_cagr_5yr",
    "eps_cagr_5yr",
]
LABELS = {
    "roe": "ROE",
    "roce": "ROCE",
    "npm": "NPM",
    "de": "D/E (inv)",
    "fcf_margin": "FCF margin",
    "pat_cagr_5yr": "PAT CAGR",
    "rev_cagr_5yr": "Rev CAGR",
    "eps_cagr_5yr": "EPS CAGR",
}


def pct_rank(s: pd.Series, lower_better: bool = False) -> pd.Series:
    """PERCENT_RANK in 0-1 where 1 = best; single member -> 0.5."""
    s = -s if lower_better else s
    n = s.notna().sum()
    if n <= 1:
        return s.where(s.isna(), 0.5)
    return (s.rank(method="min") - 1) / (n - 1)


def minmax(s: pd.Series) -> pd.Series:
    """Scale a column to 0-100 within its group; invert lower-is-better metrics."""
    lo, hi = s.min(), s.max()
    out = (
        pd.Series(50.0, index=s.index)
        if pd.isna(lo) or hi == lo
        else (s - lo) / (hi - lo) * 100
    )
    return 100 - out if s.name in LOWER_BETTER else out


def peer_percentiles(snap: pd.DataFrame, pg: pd.DataFrame) -> pd.DataFrame:
    """Long table: company_id, peer_group, metric, value, percentile_rank, year."""
    rows = []
    for group, members in pg.groupby("peer_group_name"):
        d = snap[snap["company_id"].isin(members["company_id"])]
        for m in METRICS:
            pr = pct_rank(d[m], m in LOWER_BETTER)
            for cid, yr, v, p in zip(d["company_id"], d["year"], d[m], pr):
                rows.append(
                    {
                        "company_id": cid,
                        "peer_group": group,
                        "metric": m,
                        "value": v,
                        "percentile_rank": p,
                        "year": yr,
                    }
                )
    return pd.DataFrame(rows)


def badges(pp: pd.DataFrame) -> pd.DataFrame:
    """Best in Class: top quartile on >=6 of 10 metrics; Watch List: bottom quartile on >=4."""
    g = pp.groupby(["peer_group", "company_id"])["percentile_rank"]
    b = pd.DataFrame(
        {
            "top_q": g.apply(lambda s: (s >= 0.75).sum()),
            "bottom_q": g.apply(lambda s: (s <= 0.25).sum()),
        }
    ).reset_index()
    b["badge"] = np.where(
        b["top_q"] >= 6, "Best in Class", np.where(b["bottom_q"] >= 4, "Watch List", "")
    )
    return b


def export_excel(pg: pd.DataFrame, pp: pd.DataFrame, bd: pd.DataFrame, path) -> None:
    """One sheet per peer group: 10 values + 10 percentile columns, colour-coded."""
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for group, members in pg.groupby("peer_group_name"):
            sub = pp[pp["peer_group"] == group]
            vals = sub.pivot(index="company_id", columns="metric", values="value")[
                METRICS
            ]
            pct = sub.pivot(
                index="company_id", columns="metric", values="percentile_rank"
            )[METRICS]
            sheet = vals.join(pct.add_suffix("_pctile"))
            sheet = sheet.join(
                bd[bd["peer_group"] == group].set_index("company_id")[["badge"]]
            )
            bench = set(members.loc[members["is_benchmark"] == 1, "company_id"])
            sheet.insert(
                0, "is_benchmark", ["Yes" if c in bench else "" for c in sheet.index]
            )
            name = re.sub(r"[\\/*?:\[\]]", "", group)[:31]
            sheet.round(3).to_excel(xw, sheet_name=name)
            ws = xw.sheets[name]
            for cell in ws[1]:
                if str(cell.value).endswith("_pctile"):
                    rng = f"{cell.column_letter}2:{cell.column_letter}{ws.max_row}"
                    ws.conditional_formatting.add(
                        rng,
                        CellIsRule(
                            operator="greaterThanOrEqual", formula=["0.75"], fill=green
                        ),
                    )
                    ws.conditional_formatting.add(
                        rng,
                        CellIsRule(
                            operator="lessThanOrEqual", formula=["0.25"], fill=red
                        ),
                    )
            for row in ws.iter_rows(min_row=2):
                if row[1].value == "Yes":
                    for cell in row:
                        cell.font = Font(bold=True)


def make_radars(snap: pd.DataFrame, pg: pd.DataFrame, folder) -> int:
    """One radar PNG per company vs its peer group (or sector if it has no group)."""
    folder.mkdir(parents=True, exist_ok=True)
    first = (
        pg.sort_values("peer_group_name")
        .drop_duplicates("company_id")
        .set_index("company_id")["peer_group_name"]
    )
    angles = np.linspace(0, 2 * np.pi, len(RADAR), endpoint=False).tolist()
    for cid, sector in zip(snap["company_id"], snap["broad_sector"]):
        if cid in first.index:
            group = first[cid]
            peers = snap[
                snap["company_id"].isin(
                    pg.loc[pg["peer_group_name"] == group, "company_id"]
                )
            ]
        else:
            group = f"{sector} sector"
            peers = snap[snap["broad_sector"] == sector]
        scaled = peers.set_index("company_id")[RADAR].apply(minmax)
        series = (
            (scaled.loc[cid].fillna(0).tolist(), cid, "#1f77b4"),
            (scaled.mean().fillna(0).tolist(), "Peer average", "#888888"),
        )
        fig, ax = plt.subplots(figsize=(5.5, 5.5), subplot_kw={"polar": True})
        for vals, label, color in series:
            ax.plot(angles + angles[:1], vals + vals[:1], label=label, color=color)
            ax.fill(angles + angles[:1], vals + vals[:1], alpha=0.15, color=color)
        ax.set_xticks(angles)
        ax.set_xticklabels([LABELS[m] for m in RADAR], fontsize=8)
        ax.set_ylim(0, 100)
        ax.set_title(f"{cid} vs {group}", fontsize=10)
        ax.legend(loc="upper right", bbox_to_anchor=(1.3, 1.1), fontsize=8)
        fig.savefig(folder / f"{cid}.png", dpi=120, bbox_inches="tight")
        plt.close(fig)
    return len(list(folder.glob("*.png")))


def main() -> None:
    """Compute percentiles, write peer_percentiles table, Excel and radar charts."""
    conn = sqlite3.connect(DB_PATH)
    snap = build_snapshot(conn)
    pg = pd.read_sql(
        "SELECT peer_group_name, company_id, is_benchmark FROM peer_groups", conn
    )
    pp = peer_percentiles(snap, pg)
    bd = badges(pp)
    pp.to_sql("peer_percentiles", conn, if_exists="replace", index=False)
    conn.commit()
    conn.close()
    OUT.mkdir(exist_ok=True)
    export_excel(pg, pp, bd, OUT / "peer_comparison.xlsx")
    n = make_radars(snap, pg, ROOT / "reports" / "radar_charts")
    print(
        "peer groups:",
        pp["peer_group"].nunique(),
        "| companies ranked:",
        pp["company_id"].nunique(),
    )
    print("peer_percentiles rows:", len(pp), "| radar PNGs:", n)
    print(
        bd[bd["badge"] != ""]
        .sort_values(["badge", "peer_group"])
        .to_string(index=False)
    )


if __name__ == "__main__":
    main()
