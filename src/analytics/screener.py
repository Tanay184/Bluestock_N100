"""Config-driven screener, composite screen score and Excel export."""

import operator
import sqlite3

import numpy as np
import pandas as pd
import yaml
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

from src.etl.loader import DB_PATH, OUT, ROOT

CONFIG = ROOT / "config" / "screener_config.yaml"
OPS = {
    ">": operator.gt,
    "<": operator.lt,
    ">=": operator.ge,
    "<=": operator.le,
    "==": operator.eq,
}
DE_POINTS = ([0, 0.5, 1, 2, 5], [100, 85, 70, 50, 0])
ICR_POINTS = ([1.5, 3, 5, 10], [0, 50, 75, 100])
RENAME = {
    "return_on_equity_pct": "roe",
    "roce_pct": "roce",
    "net_profit_margin_pct": "npm",
    "operating_profit_margin_pct": "opm",
    "debt_to_equity": "de",
    "interest_coverage": "icr",
    "free_cash_flow_cr": "fcf",
    "fcf_margin_pct": "fcf_margin",
    "revenue_cagr_3yr": "rev_cagr_3yr",
    "revenue_cagr_5yr": "rev_cagr_5yr",
    "dividend_payout_ratio_pct": "payout",
    "cfo_pat_ratio": "cfo_pat",
}
OUT_COLS = [
    "company_id",
    "company_name",
    "broad_sector",
    "year",
    "screen_score",
    "sector_rank",
    "roe",
    "roce",
    "npm",
    "opm",
    "de",
    "icr",
    "fcf",
    "fcf_margin",
    "fcf_yield",
    "rev_cagr_3yr",
    "rev_cagr_5yr",
    "pat_cagr_5yr",
    "eps_cagr_5yr",
    "cfo_pat",
    "pe",
    "pb",
    "ev_ebitda",
    "div_yield",
    "payout",
    "sales",
]
# column: (good threshold, bad threshold, higher_is_better)
COLOUR_RULES = {
    "roe": (15, 8, True),
    "roce": (15, 8, True),
    "npm": (10, 3, True),
    "de": (1, 2, False),
    "fcf": (0, 0, True),
    "rev_cagr_5yr": (10, 3, True),
    "pat_cagr_5yr": (15, 5, True),
}


def load_config() -> dict:
    """Read screener_config.yaml."""
    return yaml.safe_load(CONFIG.read_text(encoding="utf-8"))


def build_snapshot(conn: sqlite3.Connection) -> pd.DataFrame:
    """Latest ratios per company joined with names, sectors, valuation and sales."""
    fr = pd.read_sql("SELECT * FROM financial_ratios", conn).sort_values(
        ["company_id", "year"]
    )
    fr["de_prev"] = fr.groupby("company_id")["debt_to_equity"].shift(1)
    last = fr.groupby("company_id").tail(1).copy()
    last["mc_year"] = last["year"].str[:4].astype(int)
    comp = pd.read_sql("SELECT id AS company_id, company_name FROM companies", conn)
    sec = pd.read_sql("SELECT company_id, sub_sector FROM sectors", conn)
    mc = pd.read_sql(
        "SELECT company_id, year AS mc_year, market_cap_crore, pe_ratio AS pe, pb_ratio AS pb, "
        "ev_ebitda, dividend_yield_pct AS div_yield FROM market_cap",
        conn,
    )
    pl = pd.read_sql("SELECT company_id, year, sales FROM profitandloss", conn)
    df = last.merge(comp, on="company_id").merge(sec, on="company_id", how="left")
    df = df.merge(mc, on=["company_id", "mc_year"], how="left").merge(
        pl, on=["company_id", "year"], how="left"
    )
    df = df.rename(columns=RENAME)
    df["fcf_yield"] = df["fcf"] / df["market_cap_crore"] * 100
    df["de_change"] = df["de"] - df["de_prev"]
    return add_screen_score(df).reset_index(drop=True)


def _scale(s: pd.Series) -> pd.Series:
    """Winsorise at P10/P90 and scale to 0-100 (missing -> 0)."""
    lo, hi = s.quantile(0.10), s.quantile(0.90)
    if pd.isna(lo) or hi == lo:
        return pd.Series(0.0, index=s.index)
    return ((s.clip(lo, hi) - lo) / (hi - lo) * 100).fillna(0)


def add_screen_score(df: pd.DataFrame) -> pd.DataFrame:
    """Weighted 0-100 score: profitability 35, cash quality 30, growth 20, leverage 15."""
    prof = (
        0.15 * _scale(df["roe"]) + 0.10 * _scale(df["roce"]) + 0.10 * _scale(df["npm"])
    )
    cash = (
        0.15 * _scale(df["fcf_margin"])
        + 0.10 * _scale(df["cfo_pat"])
        + 0.05 * (df["fcf"] > 0) * 100
    )
    growth = 0.10 * _scale(df["rev_cagr_5yr"]) + 0.10 * _scale(df["pat_cagr_5yr"])
    de_s = pd.Series(np.interp(df["de"].fillna(5), *DE_POINTS), index=df.index)
    de_s = de_s.where(df["broad_sector"] != "Financials", 50.0)
    icr_s = pd.Series(np.interp(df["icr"].fillna(0), *ICR_POINTS), index=df.index)
    icr_s = icr_s.where(df["icr_label"] != "Debt Free", 100.0)
    df["screen_score"] = (prof + cash + growth + 0.10 * de_s + 0.05 * icr_s).round(1)
    df["sector_rank"] = df.groupby("broad_sector")["screen_score"].rank(
        ascending=False, method="min"
    )
    return df


def apply_filters(df: pd.DataFrame, filters: list[dict]) -> pd.DataFrame:
    """Keep rows passing every filter; missing values fail."""
    mask = pd.Series(True, index=df.index)
    for f in filters:
        mask &= OPS[f["op"]](df[f["metric"]], f["value"])
    return df[mask]


def run_preset(df: pd.DataFrame, cfg: dict, name: str) -> pd.DataFrame:
    """Run one preset, ranked by its rank_by column."""
    p = cfg["presets"][name]
    return (
        apply_filters(df, p["filters"])
        .sort_values(p["rank_by"], ascending=False)
        .head(cfg.get("top_n", 50))
    )


def export_excel(df: pd.DataFrame, cfg: dict, path) -> None:
    """Write one sheet per preset with threshold colour-coding."""
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name in cfg["presets"]:
            run_preset(df, cfg, name)[OUT_COLS].round(2).to_excel(
                xw, sheet_name=name, index=False
            )
            ws = xw.sheets[name]
            col_of = {c.value: c.column_letter for c in ws[1]}
            for col, (good, bad, hi) in COLOUR_RULES.items():
                rng = f"{col_of[col]}2:{col_of[col]}{max(ws.max_row, 2)}"
                g_op, b_op = (
                    ("greaterThanOrEqual", "lessThan")
                    if hi
                    else ("lessThanOrEqual", "greaterThan")
                )
                ws.conditional_formatting.add(
                    rng, CellIsRule(operator=g_op, formula=[str(good)], fill=green)
                )
                ws.conditional_formatting.add(
                    rng, CellIsRule(operator=b_op, formula=[str(bad)], fill=red)
                )


def main() -> None:
    """Build the snapshot, run all presets and export screener_output.xlsx."""
    conn = sqlite3.connect(DB_PATH)
    df = build_snapshot(conn)
    conn.close()
    cfg = load_config()
    OUT.mkdir(exist_ok=True)
    export_excel(df, cfg, OUT / "screener_output.xlsx")
    print("Snapshot companies:", len(df))
    for name, p in cfg["presets"].items():
        print(f"{name:22}{len(apply_filters(df, p['filters']))} companies")


if __name__ == "__main__":
    main()
